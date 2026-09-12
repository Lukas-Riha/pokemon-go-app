"""Kontrola vzorců v Pokemon_GO_Tracker.xlsx bez Excelu/LibreOffice.

Neumí vzorce spočítat (to by chtělo recalc engine), ale odchytí realistické
poruchy: nevyvážené závorky/uvozovky, odkaz na neexistující list, funkci mimo
povolenou (Google Sheets kompatibilní) sadu, `_xlfn` prefix a rozjeté párování
řádků mezi listy Import a Doporučení.
"""
import re
import sys
from pathlib import Path

import openpyxl

XLSX = Path(__file__).resolve().parent.parent / "xlsx-tracker" / "Pokemon_GO_Tracker.xlsx"

# Záměrně jen funkce, které umí i Google Sheets a starší Excel/LibreOffice.
ALLOWED = {"IF", "AND", "OR", "NOT", "COUNTIF", "COUNTIFS", "INDEX", "MATCH"}

wb = openpyxl.load_workbook(XLSX)
sheets = set(wb.sheetnames)
errors = []
total = 0

func_re = re.compile(r"([A-Z_][A-Z0-9_.]*)\s*\(")
ref_re = re.compile(r"([A-Za-zÀ-ž_][A-Za-zÀ-ž0-9_ ]*)!")


def balanced(f):
    depth = 0
    in_str = False
    i = 0
    while i < len(f):
        ch = f[i]
        if ch == '"':
            if in_str and i + 1 < len(f) and f[i + 1] == '"':
                i += 2
                continue
            in_str = not in_str
        elif not in_str:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth < 0:
                    return False, "závorka navíc"
        i += 1
    if in_str:
        return False, "neuzavřené uvozovky"
    return (depth == 0), ("chybí %d uzavíracích závorek" % depth if depth else "")


for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or not v.startswith("="):
                continue
            total += 1
            where = f"{ws.title}!{cell.coordinate}"

            ok, why = balanced(v)
            if not ok:
                errors.append(f"{where}: {why}")

            if "_xlfn" in v:
                errors.append(f"{where}: obsahuje _xlfn (nekompatibilní s Google Sheets)")

            for fn in func_re.findall(v):
                if fn not in ALLOWED:
                    errors.append(f"{where}: nepovolená funkce {fn}()")

            for ref in ref_re.findall(v):
                ref = ref.strip()
                if ref and ref not in sheets:
                    errors.append(f"{where}: odkaz na neexistující list „{ref}“")

# párování řádků Import <-> Doporučení (řádek N musí číst řádek N)
doc = wb["Doporučení"]
mismatch = 0
for r in range(5, 505):
    f = doc.cell(row=r, column=1).value or ""
    if f"Import!A{r}" not in f:
        mismatch += 1
if mismatch:
    errors.append(f"Doporučení: {mismatch} řádků neodkazuje na stejný řádek Importu")

print(f"vzorců zkontrolováno: {total}")
print(f"listy: {', '.join(wb.sheetnames)}")
if errors:
    print(f"CHYB: {len(errors)}")
    for e in errors[:20]:
        print("  -", e)
    sys.exit(1)
print("vše v pořádku (0 chyb)")
